"""Сборка двух Word-руководств из актуального кода и скриншотов.

Требует python-docx из документного окружения; к запуску игры не относится.
"""
from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path
import re
import textwrap

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from tools.guide_content import METHOD_NOTES, MODULE_INTROS, QUICK_QA, SHORT_SPEECH

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'submission'
SHOTS = ROOT / 'artifacts'
WIDTH = 7.0
SOURCES = {path.stem: path.read_text(encoding='utf-8') for path in (ROOT / 'dormgame').glob('*.py')}


def functions_in(module: str) -> dict[str, ast.FunctionDef]:
    found = {}

    def visit(node, scope=''):
        for child in ast.iter_child_nodes(node):
            if isinstance(child, (ast.ClassDef, ast.FunctionDef)):
                name = f'{scope}.{child.name}' if scope else child.name
                if isinstance(child, ast.FunctionDef):
                    found[name] = child
                visit(child, name)
            else:
                visit(child, scope)

    visit(ast.parse(SOURCES[module]))
    return found


INDEX = {module: functions_in(module) for module in SOURCES}


def configure(doc: Document, subject: str):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
    sec.header_distance = sec.footer_distance = Inches(.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
        style = doc.styles[name]
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    doc.styles['Title'].font.size = Pt(28)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].paragraph_format.space_after = Pt(15)
    doc.styles['Subtitle'].font.size = Pt(14)
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(12)
    for name, size in [('Heading 1', 20), ('Heading 2', 15), ('Heading 3', 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(8)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].paragraph_format.space_before = Pt(3)
    doc.styles['Caption'].paragraph_format.space_after = Pt(10)
    for name, size in [('Code', 9), ('Method name', 10), ('Reference body', 10.5)]:
        style = doc.styles.add_style(name, 1)
        style.font.name = 'Consolas' if name != 'Reference body' else 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.0 if name == 'Code' else 1.04
        style.paragraph_format.space_after = Pt(0 if name == 'Code' else 6)
    doc.styles['Method name'].font.bold = True
    doc.styles['Method name'].paragraph_format.space_before = Pt(9)
    doc.styles['Method name'].paragraph_format.keep_with_next = True
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = foot.add_run('byVanoGame    ')
    run.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    doc.core_properties.title = subject
    doc.core_properties.author = 'byVanoGame'
    doc.core_properties.subject = 'Подготовка к защите учебной игры на Python и Pygame'
    doc.core_properties.keywords = 'Python Pygame AABB A* подготовка защита'


def p(doc, text, style=None, bold_lead=False):
    para = doc.add_paragraph(style=style)
    if bold_lead and ': ' in text:
        first, rest = text.split(': ', 1)
        para.add_run(first + ': ').bold = True
        para.add_run(rest)
    else:
        para.add_run(text)
    return para


def page(doc, title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn('w:' + key), value)
        borders.append(element)
    t._tbl.tblPr.append(borders)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for row_index, values in enumerate([headers, *rows]):
        cells = t.rows[0].cells if row_index == 0 else t.add_row().cells
        tr_pr = cells[0]._tc.getparent().get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '243A50' if row_index == 0 else ('F0F4F7' if row_index % 2 else 'FFFFFF'))
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for direction in ('top', 'bottom', 'left', 'right'):
                margin = OxmlElement('w:' + direction)
                margin.set(qn('w:w'), '75' if direction in ('top', 'bottom') else '95')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            tcpr.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.02
            if i == 0 and widths[i] <= 1.0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(value))
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string('FFFFFF' if row_index == 0 else '000000')
            run.bold = row_index == 0
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = .5
    return t


def figure(doc, filename, caption, width=WIDTH):
    para = doc.add_paragraph()
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_after = Pt(2)
    shape = para.add_run().add_picture(str(SHOTS / filename), width=Inches(width))
    shape._inline.docPr.set('descr', caption)
    p(doc, caption, 'Caption')


def code_lines(doc, module, first, last):
    p(doc, f'dormgame/{module}.py — строки {first}–{last}', 'Caption')
    lines = SOURCES[module].splitlines()
    for number in range(first, last + 1):
        text = lines[number - 1].replace('\t', '    ')
        wrapped = textwrap.wrap(text, width=93, replace_whitespace=False,
                                drop_whitespace=False, break_long_words=False,
                                break_on_hyphens=False) or ['']
        for index, part in enumerate(wrapped):
            para = p(doc, f'{number:>3}  {part}' if index == 0 else f'     {part}', 'Code')
            para.paragraph_format.keep_with_next = number < last or index < len(wrapped) - 1
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def method_code(doc, module, name, offset=0, count=None):
    node = INDEX[module][name]
    first = node.lineno + offset
    if offset == 0 and node.decorator_list:
        first = min(decorator.lineno for decorator in node.decorator_list)
    last = node.end_lineno if count is None else min(node.end_lineno, first + count - 1)
    code_lines(doc, module, first, last)


def matching_code(doc, module, needle, before=0, count=12):
    lines = SOURCES[module].splitlines()
    start = next(i + 1 for i, line in enumerate(lines) if needle in line) - before
    code_lines(doc, module, start, min(len(lines), start + count - 1))


def command(doc, text):
    para = p(doc, text, 'Code')
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(8)
    return para


def handbook():
    doc = Document()
    configure(doc, 'Подготовка к защите игры Общага')
    doc.add_paragraph('Подготовка к защите игры Общага', 'Title')
    doc.add_paragraph('Руководство для начинающего программиста', 'Subtitle')
    p(doc, 'byVanoGame   |   Python 3.12   |   Pygame 2.6.1   |   Версия от 9 сентября 2026 года')
    figure(doc, 'menu.png', 'Рисунок 1. Главное меню актуальной версии с подписью byVanoGame.')
    p(doc, 'Это руководство помогает подготовиться к защите игры «Общага: пять минут до выселения». '
      'Здесь объясняется, как запустить проект, что показать преподавателю и как читать его код. '
      'Предполагается, что читатель почти не программировал.')
    p(doc, 'Для первого знакомства прочитайте основные главы по порядку. Справочник в конце содержит '
      'каждый явно объявленный метод и функцию игрового пакета: назначение, сигнатуру, файл и номера строк. '
      'Короткая речь и быстрые ответы вынесены во второй Word-файл.')

    page(doc, '1 Как подготовиться за один вечер')
    p(doc, 'Цель подготовки — уметь своими словами связать действие в игре с конкретным местом в коде. '
      'Например: «Нажал пробел → появилась команда → физика задала скорость вверх → Pygame нарисовал новую позицию».')
    table(doc, ['Время', 'Что сделать', 'Что должно получиться'], [
        ('10 минут', 'Запустить игру и прочитать управление', 'Двигаться, прыгать, открыть паузу и вернуться'),
        ('15 минут', 'Прочитать главы о структуре, времени и физике', 'Объяснить модель, координаты, шаг и столкновение'),
        ('15 минут', 'Разобрать Enemy и пример A*', 'Назвать наследников и объяснить g, h и f'),
        ('10 минут', 'Найти в коде пять мест из сценария показа', 'Открыть нужный файл без долгого поиска'),
        ('10 минут', 'Произнести короткую речь и ответить без подсказки', 'Уложиться примерно в две минуты')
    ], [.9, 3.1, 3.0])
    doc.add_heading('Что выучить обязательно', 2)
    p(doc, 'Где лежат правила; почему Pygame отделён от модели; зачем нужен фиксированный шаг; '
      'как работают прямоугольные коллизии; где наследование; как A* выбирает следующую клетку. '
      'Не требуется запоминать все координаты карт и команды рисования растений.')
    doc.add_heading('Как искать информацию в Word', 2)
    p(doc, 'Нажмите Ctrl+F и ищите имя метода, например update_player или astar. '
      'В области навигации Word можно переключиться на заголовки. Номера слева от кода '
      'относятся к строкам исходного Python-файла, а не к страницам руководства.')
    p(doc, 'Речь — опора для выступления. Перед защитой обязательно повторите действия в игре '
      'и объяснения своими словами: заученный текст сам по себе не помогает отвечать на уточнения.')

    page(doc, '2 Запуск и файлы для сдачи')
    p(doc, 'В этой рабочей папке уже есть подготовленное окружение. Самый простой запуск — '
      'дважды щёлкнуть start.bat. Альтернатива: открыть PowerShell в папке проекта и выполнить команду.')
    command(doc, '.\\.venv\\Scripts\\python.exe main.py')
    p(doc, 'Для переноса нужны main.py, start.bat, requirements.txt, папки dormgame, data, assets, '
      'tests и tools, а также документация. Папку .venv на другой компьютер переносить не нужно: '
      'она связана с местным Python.')
    doc.add_heading('Подготовка другого компьютера', 2)
    p(doc, 'Установите Python 3.12. Затем из папки проекта создайте окружение и установите единственную '
      'игровую зависимость. Интернет нужен на этапе установки Pygame, во время игры он не нужен.')
    for line in ['py -3.12 -m venv .venv', '.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt', '.\\.venv\\Scripts\\python.exe main.py']:
        command(doc, line)
    doc.add_heading('Если не запускается', 2)
    table(doc, ['Симптом', 'Что проверить'], [
        ('Открывается Microsoft Store', 'Используйте .venv\\Scripts\\python.exe или установленный py -3.12, а не ярлык python.'),
        ('No module named pygame', 'Установите requirements.txt именно интерпретатором из .venv.'),
        ('No module named dormgame', 'Откройте терминал в корне проекта. Проверочные сценарии запускайте через -m tools...'),
        ('Нет звука', 'Проверьте M и настройки. Без аудиоустройства игра всё равно должна работать.')
    ], [2.0, 5.0])
    p(doc, 'Скрипт build_defense_docs нужен только для повторной сборки Word. Для запуска игры '
      'устанавливать python-docx или офисные библиотеки не требуется.')

    page(doc, '3 Как играть и что изменилось')
    figure(doc, 'level1-start.png', 'Рисунок 2. Первый этаж. Наверху статистика, снизу подсказки, слева необязательный бонус.')
    table(doc, ['Клавиша', 'Действие'], [('A D или стрелки', 'Бежать'), ('Пробел', 'Прыгать; отпускание раньше даёт более низкий прыжок'),
        ('E', 'Читать записку и открыть выход'), ('R', 'Вернуться к контрольной точке'), ('Esc', 'Пауза'),
        ('F3', 'Диагностика физики и A*'), ('M и F11', 'Звук и полный экран')], [1.55, 5.45])
    p(doc, 'Выход требует все страницы текущего этажа. На трёх уровнях их 4, 4 и 5. '
      'Дополнительно на каждом старте можно пройти немного влево: там билет, кружка или флешка. '
      'Сувениры не закрывают выход и дают дополнительную строку в финале.')

    page(doc, '4 Минимум Python для чтения проекта')
    table(doc, ['Запись', 'Простой смысл'], [
        ('x = 10', 'Сохранить число 10 под именем x. Знак = присваивает, а == сравнивает.'),
        ('if условие', 'Выполнить вложенные строки, только если условие истинно.'),
        ('for item in items', 'По очереди обработать предметы списка.'),
        ('def и return', 'Объявить функцию и вернуть её результат вызывающему коду.'),
        ('class и self', 'Описать тип объекта; self означает конкретный экземпляр.'),
        ('True False None', 'Да, нет и отсутствие значения.'),
        ('list dict set tuple', 'Список, словарь ключ → значение, множество без повторов и кортеж.'),
        ('float int bool str', 'Дробное число, целое число, логическое значение и строка текста.'),
        ('_имя', 'Соглашение о внутреннем методе. Это не запрет доступа на уровне языка.')
    ], [1.8, 5.2])
    doc.add_heading('Как читать реальный короткий метод', 2)
    method_code(doc, 'geometry', 'AABB.right')
    p(doc, 'def начинает описание метода. self — этот прямоугольник. self.x — его левая координата, '
      'self.w — ширина. return возвращает сумму. Подсказка -> float сообщает ожидаемый тип результата, '
      'но сама ничего не вычисляет.')
    p(doc, 'Над методом в исходнике стоит @property: это позволяет писать box.right, а не box.right(). '
      'Отступы в Python задают вложенность. Удаление отступа может изменить смысл программы.')

    page(doc, '5 Объекты и данные без сложных слов')
    p(doc, 'Класс можно понимать как описание одинаково устроенных объектов. Объект — конкретный '
      'экземпляр: один герой, один охранник, одна страница. У каждого свои поля. '
      'Метод — действие, связанное с объектом.')
    matching_code(doc, 'commands', '@dataclass', count=9)
    p(doc, 'Commands — снимок ввода на один шаг. move равен −1, 0 или 1. jump_pressed обозначает новое '
      'нажатие, jump_held — текущее удержание. Это разные данные: удержание влияет на высоту, '
      'а новый прыжок разрешается только по новому нажатию.')
    matching_code(doc, 'model', 'class GameEvent', count=6)
    p(doc, 'GameEvent сообщает о результате шага: например, page, died или souvenir. '
      'Представление получает координаты события и показывает эффект. @dataclass автоматически '
      'создаёт обычный конструктор данных; frozen=True запрещает переустановку полей экземпляра.')
    table(doc, ['Класс данных', 'Что хранит'], [('AABB', 'x, y, ширину w и высоту h'), ('Body и PlayerBody', 'Прямоугольник, скорости, опору; у героя ещё таймеры прыжка'),
        ('Item', 'ID предмета, прямоугольник, текст и флаги collected и active'), ('CollisionResult', 'Были ли контакты по X и Y, приземление и удар о потолок'),
        ('SearchResult', 'Путь и число раскрытых узлов'), ('Particle', 'Положение, скорость, оставшуюся жизнь, цвет и размер частицы')], [1.8, 5.2])

    page(doc, '6 Карта проекта и разделение ответственности')
    table(doc, ['Файл', 'За что отвечает'], [
        ('main.py и app.py', 'Запуск окна, события устройств, меню и смена экранов'),
        ('model.py', 'Состояние сеанса, предметы, смерти, контрольные точки и выход'),
        ('physics.py', 'Движение, опора и столкновения'),
        ('geometry.py', 'Дробные прямоугольники и пересечения'),
        ('timing.py и commands.py', 'Фиксированный шаг, буфер ввода и команды'),
        ('enemies.py', 'Охранник, дрон и общий Enemy'),
        ('pathfinding.py', 'Навигационная сетка и A*'),
        ('presentation.py и audio.py', 'Рисование, камера, эффекты и звук'),
        ('levels.py и data', 'Чтение, проверка и ручные описания карт'),
        ('config.py и tests', 'Параметры и автоматические проверки')
    ], [2.25, 4.75])
    doc.add_heading('Путь одного нажатия', 2)
    p(doc, 'Клавиатура → Application.handle_event → InputBuffer → Commands → '
      'GameSession.update → update_player → новые координаты → Renderer.draw_game.')
    p(doc, 'В обратную сторону модель отдаёт GameEvent. Например, collected становится True, '
      'модель выдаёт page, затем представление добавляет частицы, а Audio играет звук.')
    p(doc, 'Разделение полезно при изменениях: можно перерисовать студента, не меняя прыжок, '
      'или проверить A* без окна. В чистых модулях нет import pygame. Все изменяемые игровые '
      'данные принадлежат объектам сеанса, а не общим глобальным переменным.')

    page(doc, '7 Как уровень получается из JSON')
    p(doc, 'JSON — текстовый формат данных. В нём нет поведения врага или формулы прыжка. '
      'Он отвечает на вопросы: где стоит платформа, где лежит страница и какой тип врага создать.')
    lines = (ROOT / 'data' / 'level1.json').read_text(encoding='utf-8').splitlines()
    p(doc, 'data/level1.json — начало реального описания уровня', 'Caption')
    for number, line in enumerate(lines[:15], 1):
        for index, part in enumerate(textwrap.wrap(line, width=91, replace_whitespace=False,
                                                 drop_whitespace=False, break_long_words=False)):
            p(doc, f'{number:>3}  {part}' if index == 0 else '     ' + part, 'Code')
    p(doc, 'rect = [x, y, w, h]. Например, [128, 394, 12, 16] означает левый верхний угол '
      'в точке 128, 394, ширину 12 и высоту 16. ID p1 отличает страницу от других предметов.')
    p(doc, 'solids — твёрдая геометрия; pages — обязательные страницы; souvenirs — необязательный набор; '
      'checkpoints — сохранения; notes — записки; hazards — опасности; exit — выход.')
    method_code(doc, 'levels', 'load_levels')
    p(doc, 'Валидатор отбрасывает неверные типы, повторные ID, отрицательные размеры, координаты '
      'за картой, появление в стене и движение платформы сквозь стену. Наличие данных ещё не '
      'доказывает проходимость: для неё есть отдельный сценарий с настоящей физикой.')

    page(doc, '8 Кадры и фиксированные шаги')
    p(doc, 'Кадр — одно обновление изображения. Шаг симуляции — одно обновление правил и физики. '
      'Они не обязаны совпадать: при 60 кадрах в секунду обычно выполняются два физических шага '
      'по 1/120 секунды на кадр.')
    method_code(doc, 'timing', 'FixedStepper.advance')
    p(doc, 'accumulator — накопленные секунды, которые ещё не потрачены на симуляцию. '
      'Пока их хватает на self.dt, вызывается update. Из накопителя вычитается один шаг. '
      'Предел max_steps не позволяет зависанию породить бесконечное догоняние.')
    p(doc, 'В паузе аккумулятор очищается. Сеанс не получает шаги, поэтому не меняются игровое '
      'время, положения врагов и защита. При продолжении старое время меню не превращается в рывок.')
    p(doc, 'Числа: FIXED_DT = 1/120 секунды; MAX_STEPS = 12; MAX_FRAME_TIME = 0,1 секунды. '
      'Сильное зависание сознательно замедляет игровое время относительно часов на стене.')

    page(doc, '9 Прыжок который удобно нажимать')
    table(doc, ['Параметр', 'Значение', 'Значение для игрока'], [
        ('RUN_SPEED', '150 px/s', 'Предельная скорость бега'), ('GRAVITY', '850 px/s²', 'Ускорение вниз'),
        ('JUMP_SPEED', '310 px/s', 'Начальная скорость вверх'), ('MAX_FALL_SPEED', '440 px/s', 'Ограничение падения'),
        ('COYOTE_TIME', '0,105 s', 'Запас после схода с края'), ('JUMP_BUFFER', '0,105 s', 'Запас нажатия до приземления')
    ], [1.8, 1.15, 4.05])
    method_code(doc, 'physics', '_jump')
    p(doc, 'Скорость вверх отрицательная, потому что y растёт вниз. _jump снимает grounded '
      'и support_id: герой уже не должен ехать вместе с платформой. Таймеры обнуляются, '
      'потому что разрешение прыжка потрачено.')
    matching_code(doc, 'physics', 'if not commands.jump_held and player.vy', count=5)
    p(doc, 'Раннее отпускание ограничивает скорость вверх величиной JUMP_CUT_SPEED. Затем '
      'гравитация меняет скорость, а скорость, умноженная на dt, даёт смещение. '
      'Если просто держать пробел, новых фронтов нет и прыжок не повторится.')
    p(doc, 'Полный прыжок в реальной симуляции поднимает героя примерно на 55,24 пикселя и '
      'переносит на 108,75 пикселя с разбега. Обычные ступени поднимаются на 32 пикселя. '
      'Это запас для честного прохождения, а не попытка использовать предельную высоту.')

    page(doc, '10 Столкновения с полом и стенами')
    p(doc, 'AABB означает прямоугольник без поворота. Для пересечения должны перекрываться '
      'и горизонтальные, и вертикальные промежутки. Координаты остаются дробными; '
      'округление происходит только при рисовании.')
    method_code(doc, 'geometry', 'AABB.intersects')
    p(doc, 'Одной проверки в конечной точке мало. Если тело быстро прошло тонкую плиту, '
      'в конце оно уже ниже неё и пересечения нет. Поэтому физика рассматривает весь '
      'промежуток движения по оси и выбирает ближайшую преграждающую поверхность.')
    method_code(doc, 'physics', 'move_body', offset=9, count=19)
    p(doc, 'Сначала X: ограничиваем движение стеной. Потом Y: ограничиваем полом или потолком. '
      'Всё смещение делится на части до 3 пикселей, а _move_axis дополнительно проверяет '
      'пересечённый промежуток. При контакте обнуляется только скорость по заблокированной оси.')
    p(doc, 'grounded заново проверяется маленьким запросом под ногами. Если платформа осталась '
      'позади, герой больше не стоит на ней. Удар головой прекращает подъём, затем действует гравитация.')

    page(doc, '11 Пространственная сетка и грузовые платформы')
    p(doc, 'Для каждого движения не нужно читать все стены карты. SpatialGrid заранее раскладывает '
      'номера стен по клеткам 24×24. Запрос проверяет только клетки рядом с движущимся телом.')
    method_code(doc, 'physics', 'SpatialGrid.query')
    p(doc, 'Одна большая стена может встретиться в нескольких клетках. indices — множество '
      'уже взятых номеров, поэтому в результате стены не дублируются. Это отбор кандидатов: '
      'точную геометрию дальше проверяет физика.')
    doc.add_heading('Перенос на движущейся опоре', 2)
    matching_code(doc, 'physics', 'carried_id = player.support_id', count=13)
    p(doc, 'Мир сначала передвигает платформу. Герой знает ID опоры и получает её dx, dy. '
      'Перенос также проверяется на столкновения. Если платформе некуда сдвинуть героя '
      'из-за стены или потолка, возвращается crushed и модель возрождает его.')
    p(doc, 'Дрон проходит через грузовую решётку по правилу игры. Для него она отсутствует '
      'и в навигации, и в физической проверке. Наземные охранники ходят по отведённым статическим участкам.')

    page(doc, '12 Наследование на примере двух врагов')
    p(doc, 'Enemy описывает общее: ID, координаты, возраст, направление и сброс. '
      'Абстрактный update_behavior требует от подкласса своё поведение. Так нельзя '
      'случайно создать «врага вообще», у которого нет алгоритма движения.')
    method_code(doc, 'enemies', 'Enemy.update')
    method_code(doc, 'enemies', 'Enemy.update_behavior')
    p(doc, 'Перед update_behavior в исходнике стоит @abstractmethod. PatrolEnemy и DroneEnemy '
      'реализуют его по-разному. Первый ходит и падает, второй летит по маршруту. '
      'Поэтому гравитация не находится в базовом Enemy.')
    matching_code(doc, 'model', 'for enemy in level.enemies:', count=3)
    p(doc, 'Это полиморфизм: общий цикл вызывает enemy.update, а вызов update_behavior '
      'попадает в реализацию конкретного класса. Фабрика build_enemy выбирает тип только '
      'при создании из JSON; цикл не перебирает условия для каждого типа врага.')
    table(doc, ['Понятие', 'Пример в проекте'], [('Наследование', 'PatrolEnemy(Enemy), DroneEnemy(Enemy)'),
        ('Абстракция', 'Обязательный update_behavior'), ('Инкапсуляция', 'Внутренние таймеры, очередь ввода и фаза платформы меняются через методы'),
        ('Композиция', 'PatrolEnemy содержит Body'), ('Полиморфизм', 'Один вызов update для разных подклассов')], [1.6, 5.4])

    page(doc, '13 A звезда на маленьком примере')
    p(doc, 'A* произносится «эй стар» или «а звезда». Он ищет кратчайший путь в графе. '
      'В этой игре вершина графа — свободная клетка, ребро — разрешённый переход в соседнюю клетку. '
      'Диагональных переходов нет.')
    table(doc, ['y / x', '0', '1', '2', '3'], [('0', 'S', '·', '#', 'G'), ('1', '·', '·', '#', '·'),
        ('2', '·', '·', '·', '·')], [1.0, 1.5, 1.5, 1.5, 1.5])
    p(doc, 'S — старт (0, 0), G — цель (3, 0), # — стена. Без стены достаточно трёх шагов вправо. '
      'Со стеной нужно спуститься на две клетки, пройти вправо и подняться: кратчайший путь имеет семь переходов.')
    table(doc, ['Символ', 'Смысл', 'Пример'], [('g', 'Сколько уже прошли от старта', 'После одного шага g = 1'),
        ('h', 'Оценка оставшегося расстояния', 'Для (1,0) до (3,0): h = 2'), ('f', 'Приоритет в очереди, g + h', 'Для этого соседа: f = 3')], [.75, 3.55, 2.7])
    method_code(doc, 'pathfinding', 'manhattan')
    p(doc, 'Очередь heapq извлекает запись с наименьшим приоритетом. h направляет поиск к цели, '
      'но не знает обо всех обходах. g хранит уже найденную реальную стоимость. '
      'Если найден более короткий подход к клетке, g обновляется.')
    p(doc, 'Манхэттенская оценка не завышает путь по четырём направлениям. Стены могут добавить '
      'шаги, но не позволяют добраться быстрее, чем сумма разниц по X и Y. '
      'Именно это помогает сохранить кратчайший результат.')

    page(doc, '14 Настоящий код поиска A звезда')
    node = INDEX['pathfinding']['astar']
    start = next(i + 1 for i, line in enumerate(SOURCES['pathfinding'].splitlines()) if 'if not grid.is_walkable(start)' in line)
    code_lines(doc, 'pathfinding', start, node.end_lineno)
    p(doc, 'Читать сверху вниз: проверка входа → очередь и словари → выбор лучшей записи → '
      'пропуск устаревшего g → восстановление при достижении цели → улучшение соседей. '
      'came_from хранит предыдущую клетку. По нему маршрут собирается с конца и разворачивается.')
    p(doc, 'В heapq нет отдельной операции уменьшения приоритета. Вместо удаления старой записи '
      'добавляется новая, а старая позже пропускается. Пустой список означает отсутствие пути; '
      'одна клетка означает совпавшие допустимые старт и цель.')

    page(doc, '15 Почему дрон обходит стену')
    figure(doc, 'astar-wall.png', 'Рисунок 3. F3 показывает маршрут под стеной, клеточную сетку, состояние дрона и раскрытые узлы.')
    p(doc, 'Навигация использует клетки 16×16, а корпус дрона имеет размер 18×14. '
      'Свободного центра недостаточно: проверяется место для всего корпуса и объём ребра '
      'между соседними центрами. Поэтому дрон не режет угол и не пролезает в слишком узкую щель.')
    table(doc, ['Состояние', 'Когда используется'], [('Патруль', 'Дрон движется между точками маршрута'),
        ('Погоня', 'Игрок приблизился на 180 пикселей или ближе'), ('Возвращение', 'Игрок оторвался на 245 пикселей; дрон возвращается к патрулю')], [1.6, 5.4])
    p(doc, 'Путь пересчитывается примерно раз в 0,45 секунды, при смене состояния или заметном '
      'смещении цели. Если клетка игрока недоступна, выбирается ближайшая точка в компоненте '
      'дрона. Компонента — область клеток, между которыми существует путь.')
    p(doc, 'Дрон замечает героя по расстоянию, в том числе через стену. Это правило обнаружения, '
      'а не разрешение проходить через геометрию. Скорость дрона 65 px/s, героя 150 px/s: можно убежать.')

    page(doc, '16 Страницы бонусы и контрольные точки')
    matching_code(doc, 'model', 'for page in level.pages:', count=10)
    p(doc, 'Проверяются два условия: предмет ещё не собран и герой пересекает его. '
      'Затем collected становится True. На следующем шаге первое условие уже ложно, '
      'поэтому счёт не увеличивается повторно. Для бонуса используется тот же простой принцип.')
    figure(doc, 'souvenir1.png', 'Рисунок 4. Подбор билета добавил один предмет в набор студента и вызвал сообщение.', width=6.45)
    p(doc, 'Контрольная точка меняет spawn. При смерти LevelState.respawn создаёт нового героя, '
      'сбрасывает врагов и платформы, но не очищает собранное. Перед переходом между этажами '
      'GameSession сохраняет количество страниц и бонусов прошлых карт.')

    page(doc, '17 От модели к картинке и звуку')
    p(doc, 'Pygame рисует внутреннюю поверхность 640×360, затем увеличивает её до окна. '
      'При другом соотношении сторон появляются полосы. Координаты мыши пересчитываются '
      'обратно, поэтому кнопки совпадают с картинкой.')
    method_code(doc, 'presentation', 'Renderer.point')
    p(doc, 'Камера — смещение взгляда. Из мировой позиции вычитается её положение. '
      'Например, мир x = 700 и камера x = 400 дают экран x = 300. Модельный герой '
      'по-прежнему остаётся в 700: камера его не перемещает.')
    figure(doc, 'ending.png', 'Рисунок 5. Финал автоматического прохождения: 13 страниц, 3 сувенира и 0 смертей.', width=6.45)
    p(doc, 'Спрайты, шрифты и свечения кэшируются. Частиц не больше 90, а при уменьшенных '
      'эффектах остаётся до 20. Audio создаёт короткие тоны и собственный луп в памяти; '
      'при отсутствии устройства звук пропускается. Настройки хранятся в папке пользователя.')

    page(doc, '18 Проверки и честные результаты')
    p(doc, 'Текущая версия прошла 75 автоматических тестов. Чистые тесты запускаются даже с -S, '
      'который отключает сторонние пакеты. Это показывает, что правила, физика и AI не требуют Pygame.')
    command(doc, '.\\.venv\\Scripts\\python.exe -S -m unittest discover -v')
    table(doc, ['Проверка', 'Что она доказывает'], [('Физика', 'Контакты с полом, потолком и стенами; тонкая плита; прыжки и движущиеся опоры'),
        ('A* против BFS', 'Длина пути совпадает с независимым алгоритмом на небольших картах'),
        ('Модель', 'Предметы не дублируются, смерть сохраняет собранное, выход требует страницы'),
        ('Маршрут', 'Обычные команды могут провести героя через три настоящих уровня'),
        ('Pygame', 'Меню, ввод, пауза, смена экранов и перезапуск работают вместе')], [1.6, 5.4])
    command(doc, '.\\.venv\\Scripts\\python.exe -m tools.verify_routes')
    command(doc, '.\\.venv\\Scripts\\python.exe -m tools.verify_pygame --native')
    p(doc, 'Прогон Pygame с бонусами завершился за 33,367 секунды игрового времени: '
      '13 страниц, 3 сувенира, 0 смертей и 20 прыжков. Это точный контроллер, а не '
      'оценка времени прохождения человеком.')
    p(doc, 'Человек не проходил игру вручную в рамках этих проверок. Сохранённые кадры '
      'просмотрены, а окно Windows запускалось. Не стоит говорить преподавателю, что '
      'автоматический прогон доказывает субъективное удобство для всех игроков.')

    page(doc, '19 Что показать преподавателю')
    p(doc, 'Держите рядом игру и редактор кода. Начните с короткой речи из второго документа. '
      'Затем покажите пять действий, каждое связано с конкретной частью реализации.')
    table(doc, ['Шаг', 'Действие', 'Объяснение'], [('1', 'Открыть первый этаж, пройти влево и подобрать билет', 'Необязательный предмет и событие модели'),
        ('2', 'Прыгнуть коротко, затем удержать пробел', 'Управляемая высота и скорость вверх'),
        ('3', 'Добраться до сохранения и нажать R', 'Новый герой в checkpoint, страницы остаются'),
        ('4', 'Открыть второй этаж и включить F3 возле стены', 'A* ищет обход для корпуса дрона'),
        ('5', 'Открыть Enemy.update и GameSession.update', 'Один контракт, разные реализации поведения')], [.5, 3.1, 3.4])
    command(doc, '.\\.venv\\Scripts\\python.exe main.py --level 2 --debug')
    p(doc, 'Быстрый запуск открывает второй этаж для демонстрации. До решётки нужно дойти '
      'по первым ступеням. Стена находится справа от неё. По F3 видны розовые коллайдеры, '
      'жёлтый путь, состояние дрона и число раскрытий.')
    p(doc, 'На сцене A* не нужно выигрывать гонку с дроном: достаточно увидеть обход и открыть '
      'паузу. По Escape таймеры останавливаются, можно спокойно объяснить рисунок. '
      'В коде найдите astar по Ctrl+F и покажите g_score, очередь и came_from.')
    p(doc, 'Если запутались в названии, объясните действие обычными словами и покажите код. '
      'Например: «Здесь хранится предыдущая клетка, чтобы потом собрать маршрут». '
      'Точность полезнее большого количества терминов.')

    for part, questions in enumerate((QUICK_QA[:9], QUICK_QA[9:]), 20):
        page(doc, f'{part} Вопросы преподавателя и ответы')
        for question, answer in questions:
            para = p(doc, question)
            para.runs[0].bold = True
            para.paragraph_format.keep_with_next = True
            p(doc, answer)

    page(doc, '22 Небольшая практика перед защитой')
    p(doc, 'Выполняйте упражнения в копии папки проекта. После изменения запустите игру, '
      'а затем тесты. Перед самой сдачей вернитесь к проверенной версии.')
    for title, text in [
        ('Объяснить один предмет', 'Откройте GameSession.update. Найдите collected = True. Объясните, почему следующая проверка не засчитает предмет снова. Затем найдите событие souvenir в Audio и Renderer.'),
        ('Предсказать изменение скорости', 'Найдите RUN_SPEED в config.py. Увеличение меняет максимальную скорость, но не делает её мгновенной: разгон задаёт ACCELERATION. Скажите, какие тесты и прыжки нужно проверить после такого изменения.'),
        ('Изменить одну реплику', 'В data/level1.json найдите text у страницы. Измените только текст. Геометрия, ID и условие выхода останутся прежними. Это пример разделения данных и поведения.'),
        ('Нарисовать путь на бумаге', 'Повторите маленькую карту из главы об A*. Подпишите старт, цель, стену и один путь из семи переходов. Объясните, почему прямой путь закрыт.'),
        ('Прочитать метод вслух', 'Выберите AABB.intersects или GameSession._die. Для каждой строки назовите входные данные, действие и результат. После этого объясните метод без чтения.')]:
        doc.add_heading(title, 2)
        p(doc, text)
    p(doc, 'Самопроверка: можете ли вы за минуту объяснить, где координаты, где рисунок, '
      'почему нужен dt и чем очередь событий отличается от очереди поиска A*? '
      'Если нет, вернитесь к соответствующей главе, а не к заучиванию всей речи.')

    page(doc, '23 Критерии оценки и ограничения')
    table(doc, ['Критерий', 'Вес', 'Что показать'], [('DRY', '2', 'Общие AABB, move_body, Enemy.update и кэши'),
        ('KISS', '2', 'Короткие ручные карты и простые структуры'), ('SOLID', '2', 'Разделение модулей и общий контракт врагов'),
        ('Качество кода', '1', 'Именованные параметры, типы и проверка данных'), ('Инкапсуляция абстракция наследование', '3', 'Enemy, внутренние таймеры и два наследника'),
        ('Сложность алгоритма', 'до 15', 'Самостоятельный A*, heapq, g-score и BFS-тесты'), ('Разделение ответственности', '10', 'Модель, физика, AI и представление'),
        ('Структура проекта', '5', 'Пакет, данные, тесты и документы'), ('Работоспособность', '3', 'Три этажа, финал и новая игра'),
        ('Удобство', '2', 'Буфер прыжка, coyote time, пауза и сохранения'), ('Сюжет оформление музыка', '6', 'Ночная общага, реплики, сувениры и локальный луп')], [2.3, .65, 4.05])
    p(doc, 'Вес критерия взят из задания. Это не обещание полученной оценки: на защите '
      'важны работающая демонстрация и понимание реализации.')
    p(doc, 'Ограничения: нет склонов, поворачиваемых коллайдеров, редактора карт и сохранения '
      'прохождения между запусками. Настройки сохраняются. Дрон видит по расстоянию и '
      'игнорирует грузовые решётки по правилам мира. A* находит кратчайший путь в своей '
      'сетке, а не произвольную идеальную траекторию в непрерывном пространстве.')

    page(doc, '24 Оценка сложности без завышенных обещаний')
    p(doc, 'V — допустимые клетки навигации, E — переходы, C — все клетки, S — статические '
      'прямоугольники, L — длина пути. Для сетки с четырьмя направлениями E не больше 4V.')
    table(doc, ['Операция', 'Время', 'Почему'], [('A*', 'O((V + E) log V)', 'Очередь приоритетов обрабатывает вершины и улучшения рёбер'),
        ('A* на этой сетке', 'O(V log V)', 'Число рёбер пропорционально числу вершин'), ('Восстановление', 'O(L)', 'Проходим по came_from и разворачиваем список'),
        ('Построение навигации', 'O(C·S + V log V)', 'Однократные проверки стен и сортировка свободных клеток'),
        ('Компоненты', 'O(V + E)', 'Обход связей графа'), ('Точка приближения', 'O(Vₖ) в худшем случае', 'Просмотр компоненты дрона, если клетка игрока недоступна'),
        ('Запрос spatial grid', 'O(Q + K) в среднем', 'Q посещённых клеток и K прочитанных ссылок, включая повторы')], [1.8, 1.7, 3.5])
    p(doc, 'Память A* — O(V + E), здесь O(V). Словарь и множество считаются имеющими '
      'среднюю постоянную стоимость отдельной операции. Но весь spatial grid-запрос '
      'не O(1): он зависит от количества затронутых клеток и ссылок.')
    p(doc, 'В движении дрона путь хранится коротким списком; удаление первой клетки стоит O(L). '
      'Это стоимость следования пути, а не самого алгоритма поиска. Для текущих небольших '
      'карт решение достаточно, а на длинных маршрутах можно использовать очередь или индекс.')

    page(doc, '25 Полный справочник функций и методов игры')
    count = sum(len(items) for items in INDEX.values())
    p(doc, f'Ниже перечислены все {count} явно объявленные функции и методы пакета dormgame, '
      'включая внутренние методы, свойства и вложенные помощники валидатора. Список извлечён '
      'из Python-кода автоматически; пояснения написаны отдельно. Сборщик проверяет, '
      'что у каждого объявления есть описание.')
    p(doc, 'Файлы config.py и commands.py содержат параметры и dataclass без собственных '
      'методов. __init__.py объявляет пакет. main.py передаёт управление app.main. '
      'Автоматически создаваемые dataclass методы не имеют строк в исходном файле '
      'и поэтому не выдаются за написанные вручную. Тестовые и служебные скрипты '
      'разобраны по назначению в главе о проверках; они не входят в игровой пакет.')
    p(doc, 'Сигнатура показывает имя и параметры. После стрелки указан ожидаемый тип результата. '
      'Если метод возвращает None, он обычно изменяет свой объект или выполняет действие. '
      'Фрагменты сигнатур ниже взяты из текущих файлов; многоточие в длинном объявлении '
      'означает только перенос записи в справочнике, а не пропуск реализации в проекте.')
    for module in MODULE_INTROS:
        doc.add_heading(f'Модуль {module}', 2)
        p(doc, MODULE_INTROS[module]).paragraph_format.keep_with_next = True
        for name, node in INDEX[module].items():
            para = p(doc, name, 'Method name')
            source_lines = SOURCES[module].splitlines()
            signature = source_lines[node.lineno - 1].strip()
            cursor = node.lineno
            while not signature.rstrip().endswith((':', '...')) and cursor < node.end_lineno:
                signature += ' ' + source_lines[cursor].strip()
                cursor += 1
            info = p(doc, f'dormgame/{module}.py  |  строки {node.lineno}–{node.end_lineno}', 'Caption')
            info.paragraph_format.keep_with_next = True
            signature_para = p(doc, signature, 'Code')
            signature_para.paragraph_format.keep_with_next = True
            p(doc, METHOD_NOTES[module][name], 'Reference body')
    doc.add_heading('Как воспроизвести материал', 2)
    p(doc, 'Скриншоты снимает tools.verify_pygame, а Word-файлы собирает tools.build_defense_docs. '
      'Справочник ссылается на эту версию исходников. После изменения кода номера строк '
      'могут сдвинуться; для новой защиты документы нужно пересобрать.')
    return doc


def speech():
    doc = Document()
    configure(doc, 'Короткая речь для защиты игры Общага')
    doc.add_paragraph('Короткая речь для защиты игры Общага', 'Title')
    p(doc, 'byVanoGame   |   Примерно полторы или две минуты спокойным темпом')
    for paragraph in SHORT_SPEECH:
        p(doc, paragraph)
    p(doc, 'Перед выступлением произнесите текст вслух и замените непривычные слова своими. '
      'Названия методов можно показать на экране. Не заявляйте о проверках или своём опыте, '
      'которые не можете подтвердить.')
    page(doc, 'Что открыть и показать после речи')
    command(doc, '.\\.venv\\Scripts\\python.exe main.py')
    command(doc, '.\\.venv\\Scripts\\python.exe main.py --level 2 --debug')
    for step in [
        '1. Главное меню. Покажите byVanoGame, управление и настройки.',
        '2. Первый этаж. Короткий и полный прыжок, страница и необязательный бонус слева.',
        '3. Контрольная точка. После её активации нажмите R: собранное останется.',
        '4. Второй этаж. Дойдите до решётки, включите F3 и покажите обход стены дроном.',
        '5. Код. Откройте Enemy.update, GameSession.update, astar и update_player.'
    ]:
        p(doc, step)
    doc.add_heading('Пять быстрых ответов', 2)
    for index in (1, 2, 5, 6, 11):
        question, answer = QUICK_QA[index]
        p(doc, question + ': ' + answer, bold_lead=True)
    page(doc, 'Последняя проверка перед сдачей')
    p(doc, 'Должны запускаться игра, второй этаж с F3 и тесты. На чужом компьютере '
      'проверьте это заранее, а не перед началом выступления.')
    command(doc, '.\\.venv\\Scripts\\python.exe -m unittest discover -v')
    p(doc, 'Проверенная версия: 75 тестов; 3 уровня; 13 страниц; 3 необязательных сувенира. '
      'Автоматическое прохождение Pygame с бонусами — 0 смертей и 20 прыжков. '
      'Ручного полного прохождения в отчёте нет.')
    for index in (7, 8, 9, 10, 12, 13, 14):
        question, answer = QUICK_QA[index]
        p(doc, question + ': ' + answer, bold_lead=True)
    p(doc, 'Если вопрос незнакомый: спокойно уточните, какую часть кода нужно объяснить. '
      'Откройте метод и расскажите, какие данные входят, что меняется и что возвращается. '
      'Не придумывайте ответ про оптимизацию или тест, которых в проекте нет.')
    return doc


def main():
    missing = [(module, name) for module, items in INDEX.items() for name in items
               if name not in METHOD_NOTES.get(module, {})]
    extra = [(module, name) for module, items in METHOD_NOTES.items() for name in items
             if name not in INDEX.get(module, {})]
    if missing or extra:
        raise ValueError(f'Покрытие справочника: отсутствуют={missing}, лишние={extra}')
    OUTPUT.mkdir(exist_ok=True)
    handbook().save(OUTPUT / 'Руководство_по_защите_byVanoGame.docx')
    speech().save(OUTPUT / 'Короткая_речь_и_шпаргалка_byVanoGame.docx')
    report = {
        'methods': sum(len(items) for items in INDEX.values()),
        'modules': {module: len(items) for module, items in INDEX.items()},
        'source_sha256': {module: hashlib.sha256(text.encode()).hexdigest() for module, text in SOURCES.items()},
        'speech_words': len(' '.join(SHORT_SPEECH).split()),
        'source_references': {module: {name: [node.lineno, node.end_lineno] for name, node in items.items()}
                              for module, items in INDEX.items()},
    }
    (ROOT / 'artifacts' / 'guide-coverage.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Созданы 2 Word-файла. Методов: {report["methods"]}. Слов в речи: {report["speech_words"]}.')


if __name__ == '__main__':
    main()
